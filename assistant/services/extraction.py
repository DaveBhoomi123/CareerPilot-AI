from pathlib import Path
from zipfile import ZipFile, BadZipFile
from django.core.exceptions import ValidationError
from pypdf import PdfReader
from docx import Document

def extract_resume(upload):
    if upload.size > 5 * 1024 * 1024:
        raise ValidationError('Please upload a file smaller than 5 MB.')
    extension = Path(upload.name).suffix.lower()
    try:
        if extension == '.txt':
            text = upload.read().decode('utf-8-sig')
        elif extension == '.pdf':
            reader = PdfReader(upload)
            if reader.is_encrypted or len(reader.pages) > 20:
                raise ValueError('Encrypted or too many pages')
            text = '\n'.join(page.extract_text() or '' for page in reader.pages)
        elif extension == '.docx':
            with ZipFile(upload) as archive:
                if sum(i.file_size for i in archive.infolist()) > 20 * 1024 * 1024:
                    raise ValueError('Expanded document too large')
            upload.seek(0)
            doc = Document(upload)
            text = '\n'.join([p.text for p in doc.paragraphs] + [cell.text for table in doc.tables for row in table.rows for cell in row.cells])
        else:
            raise ValidationError('Supported formats: PDF, DOCX and UTF-8 TXT.')
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError('Could not read this document. Use an unencrypted text-based PDF, DOCX or UTF-8 TXT.') from exc
    text = text.replace('\x00', '').strip()
    if len(text) < 50:
        raise ValidationError('At least 50 characters of readable text are required. Scanned PDFs need OCR first.')
    if len(text) > 40000:
        raise ValidationError('Resume text must be under 40,000 characters.')
    return text
